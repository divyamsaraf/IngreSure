"""Item 11: INFOODS AnFooD SourceAdapter registration + emission."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from core.knowledge.ike2.etl.adapt import map_record
from core.knowledge.ike2.etl.adapters import SourceAdapter, emit_layer1_records
from core.knowledge.ike2.etl.adapters.infoods_anfood import InfoodsAnfoodAdapter
from core.knowledge.ike2.etl.bulk_inject import inject
from core.knowledge.ike2.etl.license_check import clear_allowlist_cache
from core.knowledge.ike2.etl.reconcile import _SOURCE_TIER, _tier
from core.knowledge.ike2.etl.sources import canonical_source, default_knowledge_state
from core.knowledge.ike2.etl.validate import validate_rows
from tests.ike2.test_etl_bulk_inject import FakeWriter

_REPO = Path(__file__).resolve().parents[2].parent


@pytest.fixture(autouse=True)
def _reload_license_allowlist():
    clear_allowlist_cache()
    yield
    clear_allowlist_cache()


def _mini_anfood_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    # Drop default sheet name noise
    default = wb.active
    wb.remove(default)

    fish = wb.create_sheet("09 Fish & Shellfish")
    fish.append(
        [
            "Food Item ID",
            "Subgroup",
            "Country, region",
            "Food name in own language",
            "Food name in English",
            "Processing",
            "Scientific name",
        ]
    )
    fish.append(
        [
            "F1",
            "Finfish",
            "Malawi",
            "mlamba",
            "African catfish farmed skinless boneless fillet boiled",
            "p",
            "Clarias gariepinus",
        ]
    )
    fish.append(
        [
            "F2",
            "Molluscs",
            "Unknown",
            "",
            "Abalone edible flesh raw",
            "r",
            "Haliotis midae",
        ]
    )

    veg = wb.create_sheet("05 Vegetables")
    veg.append(
        [
            "Food Item ID",
            "Country, region",
            "Foodname in own language",
            "Foodname in English",
            "Processing",
            "Scientific name",
        ]
    )
    veg.append(
        ["V1", "India", "palak", "Spinach raw", "r", "Spinacia oleracea"]
    )
    veg.append(
        ["V2", "India", "", "Spinach boiled", "p", "Spinacia oleracea"]
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_infoods_adapter_is_source_adapter():
    assert isinstance(InfoodsAnfoodAdapter(), SourceAdapter)


def test_infoods_license_ok_for_cc_by():
    assert InfoodsAnfoodAdapter().license_ok() is True


def test_infoods_registered_auto_classified_tier3():
    assert canonical_source("infoods_anfood") == "infoods"
    assert default_knowledge_state("infoods") == "AUTO_CLASSIFIED"
    assert _tier("infoods") == 3
    assert "infoods" in _SOURCE_TIER


def test_infoods_emit_head_canonical_and_aliases(tmp_path):
    xlsx = _mini_anfood_xlsx(tmp_path / "AnFooD2.0.xlsx")
    rows = emit_layer1_records(InfoodsAnfoodAdapter(xlsx_path=xlsx))
    by = {r["canonical_name"]: r for r in rows}

    catfish = by["african catfish"]
    assert "African catfish farmed skinless boneless fillet boiled" in catfish["aliases"]
    assert "Clarias gariepinus" in catfish["aliases"]
    assert "mlamba" in catfish["aliases"]
    assert catfish["animal_origin"] is True
    assert catfish["fish_source"] is True
    assert catfish.get("plant_origin") is not True

    abalone = by["abalone"]
    assert abalone["shellfish_source"] is True
    assert abalone.get("fish_source") is not True

    spinach = by["spinach"]
    assert spinach["plant_origin"] is True
    aliases_l = {a.lower() for a in spinach["aliases"]}
    assert "spinach raw" in aliases_l
    assert "spinach boiled" in aliases_l
    assert "palak" in aliases_l


def test_infoods_records_validate_and_inject_auto_classified(tmp_path):
    xlsx = _mini_anfood_xlsx(tmp_path / "AnFooD2.0.xlsx")
    raw_rows = emit_layer1_records(InfoodsAnfoodAdapter(xlsx_path=xlsx))
    mapped = [
        map_record(r, "infoods", default_knowledge_state("infoods"))[0]
        for r in raw_rows
    ]
    ok, rejects = validate_rows(mapped)
    assert rejects == []
    assert all(r["knowledge_state"] == "AUTO_CLASSIFIED" for r in ok)
    # asserted_flags from map_record: deliberate keys only
    sample = next(r for r in ok if r["canonical_name"] == "african catfish")
    assert "animal_origin" in sample["asserted_flags"]
    assert "fish_source" in sample["asserted_flags"]
    assert "synthetic" not in sample["asserted_flags"]

    w = FakeWriter()
    stats = inject(raw_rows, "infoods", w)
    assert stats.rejected == 0
    assert stats.inserted == len(raw_rows)
    assert w.get_group("african catfish")["knowledge_state"] == "AUTO_CLASSIFIED"


def test_infoods_license_deny_blocks_emit(tmp_path, monkeypatch):
    xlsx = _mini_anfood_xlsx(tmp_path / "AnFooD2.0.xlsx")
    # Wrong SPDX for this source → allowlist deny
    ad = InfoodsAnfoodAdapter(xlsx_path=xlsx, license_spdx="CC0-1.0")
    assert emit_layer1_records(ad) == []
