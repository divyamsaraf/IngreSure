"""INFOODS AnFooD2.0 SourceAdapter (Item 11).

Composition-table feed: head commodity is the canonical name; the full
analytical English string, scientific name, and local name are aliases.
Flags are inferred from workbook sheet membership (and fish subgroup).
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from core.knowledge.ike2.etl.license_check import license_check
from core.normalization.normalizer import normalize_ingredient_key

_REPO_ROOT = Path(__file__).resolve().parents[6]
_DEFAULT_XLSX = _REPO_ROOT / "data" / "raw" / "infoods" / "AnFooD2.0.xlsx"

# Processing / prep tokens stripped when forming the head commodity.
_PROC = re.compile(
    r"\b("
    r"raw|boiled|fried|grilled|baked|roasted|dried|cooked|smoked|canned|"
    r"microwaved|blanched|farmed|wild|fillet|fillets|flesh|muscle|loin|"
    r"skinless|boneless|bonless|powder|flour|edible|whole|fresh|"
    r"oven|sunflower\s+oil|vegetable\s+oil"
    r")\b",
    re.I,
)

_PLANT_SHEET_PREFIXES = (
    "01 ",
    "02 ",
    "03 ",
    "04 ",
    "05 ",
    "06 ",
    "11 ",
)
_ANIMAL_SHEET_PREFIXES = (
    "07 ",
    "08 ",
    "09 ",
    "10 ",
)

_SHELLFISH_SUBGROUPS = frozenset({
    "molluscs",
    "mollusc",
    "mollusk",
    "mollusks",
    "crustacean",
    "crustaceans",
    "bivalve",
    "bivalves",
    "cephalopod",
    "cephalopods",
    "gastropod",
    "gastropods",
    "shellfish",
})

# Name lexicon for sheet-09 rows whose subgroup is missing/mis-coded (e.g. scallop).
_SHELLFISH_NAME = re.compile(
    r"\b(prawns?|shrimps?|crabs?|lobsters?|crawfish|crayfish|crawdads?|"
    r"langoustines?|krill|yabb(?:y|ies)|mussels?|oysters?|clams?|scallops?|"
    r"squids?|octopuses?|octopus|cuttlefish|shellfish|molluscs?|mollusks?|"
    r"crustaceans?|abalone|whelk|cockle|periwinkle)\b",
    re.I,
)


def head_commodity(english: str) -> str:
    """Primary canonical form: commodity head, not the analytical full string."""
    s = (english or "").strip()
    if not s:
        return ""
    s = re.split(r"[;\n]", s, maxsplit=1)[0]
    # Drop trailing ", raw" / leading descriptives after first comma clause keep head
    if "," in s:
        head, rest = s.split(",", 1)
        # If the part before comma is short, use it; else strip proc from whole
        if len(head.split()) <= 6:
            s = head
        else:
            s = head + " " + rest
    s = _PROC.sub(" ", s)
    s = re.sub(r"\s+", " ", s).strip(" -_,")
    return normalize_ingredient_key(s)


def flags_for_sheet(
    sheet: str, subgroup: str | None, *, english: str = ""
) -> dict[str, Any]:
    """Infer origin/allergen flags from AnFooD sheet (+ fish subgroup).

    Emit only positive decisions (plus uncertainty). Explicit false keys are the
    asserted_flags pollution signature when adapt uses key-presence.
    """
    name = (sheet or "").strip()
    flags: dict[str, Any] = {
        "uncertainty_flags": ["infoods_anfood_inferred"],
    }
    if any(name.startswith(p) for p in _PLANT_SHEET_PREFIXES):
        flags["plant_origin"] = True
    if name.startswith("07 "):
        flags["animal_origin"] = True
    if name.startswith("08 "):
        flags["animal_origin"] = True
        flags["egg_source"] = True
    if name.startswith("10 "):
        flags["animal_origin"] = True
        flags["dairy_source"] = True
    if name.startswith("09 "):
        flags["animal_origin"] = True
        sub = (subgroup or "").strip().lower()
        # Subgroup token or commodity name — sheet 09 defaults to fish when
        # neither says shellfish (american sea scallop had empty/wrong subgroup).
        if sub in _SHELLFISH_SUBGROUPS or _SHELLFISH_NAME.search(english or ""):
            flags["shellfish_source"] = True
        else:
            flags["fish_source"] = True
    return flags


def _region_token(country: str) -> str | None:
    """Best-effort region; prefer ISO-ish when obvious, else None (Global)."""
    c = (country or "").strip()
    if not c or c.lower() in ("unknown", "n/a", "na"):
        return None
    # Take leading country word before comma
    lead = re.split(r"[,;(]", c, maxsplit=1)[0].strip()
    aliases = {
        "usa": "US",
        "united states": "US",
        "u.s.": "US",
        "u.s.a.": "US",
        "uk": "UK",
        "united kingdom": "UK",
        "india": "IN",
        "canada": "CA",
        "australia": "AU",
        "new zealand": "NZ",
        "germany": "DE",
        "france": "FR",
        "italy": "IT",
        "brazil": "BR",
        "nigeria": "NG",
        "malawi": "MW",
        "tanzania": "TZ",
        "iceland": "IS",
        "ireland": "IE",
        "mexico": "MX",
    }
    key = lead.lower()
    if key in aliases:
        return aliases[key]
    # Unknown free-text country — do not invent a region stamp
    return None


def normalize_anfood_row(
    *,
    sheet: str,
    english: str,
    local: str,
    scientific: str,
    subgroup: str | None,
    country: str,
    food_id: str,
) -> dict[str, Any] | None:
    """One AnFooD row → Layer-1 dict (head commodity as canonical)."""
    full = (english or "").strip()
    if not full:
        return None
    canon = head_commodity(full)
    if not canon:
        return None

    flags = flags_for_sheet(sheet, subgroup, english=full)
    aliases: list[str] = []
    seen: set[str] = {canon}

    def _add(alias: str) -> None:
        a = (alias or "").strip()
        if not a:
            return
        key = normalize_ingredient_key(a)
        if not key or key in seen:
            return
        aliases.append(a)
        seen.add(key)

    _add(full)
    _add(scientific)
    _add(local)
    if food_id:
        _add(f"INFOODS:{food_id.strip()}")

    region = _region_token(country)
    regions = [region] if region else ["Global"]

    rec: dict[str, Any] = {
        "canonical_name": canon,
        "aliases": aliases,
        "source": "infoods",
        "regions": regions,
        **flags,
    }
    if region:
        rec["region"] = region
    if flags.get("fish_source"):
        rec["animal_species"] = "fish"
    elif flags.get("shellfish_source"):
        rec["animal_species"] = "shellfish"
    return rec


def merge_anfood_records(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Dedupe by normalized canonical; union aliases; OR sticky true flags."""
    merged: dict[str, dict[str, Any]] = {}
    bool_or = (
        "plant_origin",
        "animal_origin",
        "fish_source",
        "shellfish_source",
        "dairy_source",
        "egg_source",
    )
    for row in rows:
        if not row:
            continue
        key = normalize_ingredient_key(row["canonical_name"])
        if key not in merged:
            merged[key] = dict(row)
            merged[key]["aliases"] = list(row.get("aliases") or [])
            continue
        dest = merged[key]
        seen = {normalize_ingredient_key(a) for a in dest.get("aliases") or []}
        seen.add(key)
        for a in row.get("aliases") or []:
            an = normalize_ingredient_key(a)
            if an and an not in seen:
                dest.setdefault("aliases", []).append(a)
                seen.add(an)
        for f in bool_or:
            if row.get(f):
                dest[f] = True
        # Prefer animal_species when newly animalish
        if row.get("animal_species") and not dest.get("animal_species"):
            dest["animal_species"] = row["animal_species"]
        # Union regions
        regs = list(dest.get("regions") or [])
        for r in row.get("regions") or []:
            if r not in regs:
                regs.append(r)
        dest["regions"] = regs
        u = list(dest.get("uncertainty_flags") or [])
        for flag in row.get("uncertainty_flags") or []:
            if flag not in u:
                u.append(flag)
        dest["uncertainty_flags"] = u
    return sorted(merged.values(), key=lambda r: r["canonical_name"])


def _col_index(header: list[str], *needles: str) -> int | None:
    normalized = [h.lower().replace(" ", "") for h in header]
    for needle in needles:
        n = needle.lower().replace(" ", "")
        for i, h in enumerate(normalized):
            if n in h:
                return i
    return None


def read_anfood_xlsx(path: Path) -> list[dict[str, Any]]:
    """Parse AnFooD2.0 workbook into merged Layer-1 records."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    raw_rows: list[dict[str, Any]] = []
    for sheet_name in wb.sheetnames:
        if sheet_name in (
            "Copyright",
            "Codes",
            "Food groups",
            "Components",
            "Bibliography",
        ):
            continue
        if "fatty" in sheet_name.lower():
            continue
        ws = wb[sheet_name]
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            continue
        header = [str(h).strip() if h is not None else "" for h in header_row]
        i_en = _col_index(header, "foodname in english", "food name in english")
        if i_en is None:
            continue
        i_local = _col_index(
            header, "foodname in own language", "food name in own language", "foodname in"
        )
        i_sci = _col_index(header, "scientific name", "species")
        i_sub = _col_index(header, "subgroup")
        i_country = _col_index(header, "country")
        i_id = _col_index(header, "food item id", "fooditemid")
        for row in it:
            def cell(idx: int | None) -> str:
                if idx is None or idx >= len(row) or row[idx] is None:
                    return ""
                return str(row[idx]).strip()

            english = cell(i_en)
            if not english:
                continue
            rec = normalize_anfood_row(
                sheet=sheet_name,
                english=english,
                local=cell(i_local),
                scientific=cell(i_sci),
                subgroup=cell(i_sub) or None,
                country=cell(i_country),
                food_id=cell(i_id),
            )
            if rec:
                raw_rows.append(rec)
    return merge_anfood_records(raw_rows)


@dataclass
class InfoodsAnfoodAdapter:
    """FAO/INFOODS AnFooD2.0 → Layer-1 via the shared license gate."""

    name: str = "infoods"
    license_spdx: str = "CC-BY-4.0"
    xlsx_path: Path = field(default_factory=lambda: _DEFAULT_XLSX)

    def license_ok(self) -> bool:
        return bool(license_check(self.name, self.license_spdx))

    def to_layer1_records(self) -> list[dict]:
        path = Path(self.xlsx_path)
        if not path.is_file():
            return []
        return read_anfood_xlsx(path)


def write_layer1_dump(
    path: Path,
    *,
    adapter: InfoodsAnfoodAdapter | None = None,
) -> list[dict]:
    """Emit Layer-1 JSON via ``emit_layer1_records``."""
    from core.knowledge.ike2.etl.adapters.protocol import emit_layer1_records

    ad = adapter or InfoodsAnfoodAdapter()
    rows = emit_layer1_records(ad)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return rows
